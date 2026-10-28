import openpyxl

wb = openpyxl.load_workbook('DOCUMENTOS BASE/COSTOS_GASTOS_SEPTIEMBRE_2026.xlsx', data_only=True)

target = 235920150
print(f"Target: {target:,.2f}")

# Let's inspect all numbers in the entire workbook
matches = []
for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                if abs(v - target) < 100:
                    matches.append((name, r, c, v))
                # Check if diff is interesting
                diff = target - v
                if abs(diff) < 1000000:
                    print(f"Close match in {name} R{r}C{c}: {v:,.2f} (diff: {diff:,.2f})")

print("Exact matches:", matches)
