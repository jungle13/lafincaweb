import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import os
import re
import urllib.request
import collections

print("=== INICIANDO CONSTRUCCIÓN DEL EXCEL DE MOVIMIENTOS HISTÓRICOS DE SEPTIEMBRE ===")

# -------------------------------------------------------------------------
# 1. CARGAR CATÁLOGO DE INSUMOS DESDE SUPABASE
# -------------------------------------------------------------------------
headers = {
    'apikey': 'sb_publishable_VhddzdHirUzYfgMe7shYDA_-i7T4xnO',
    'Authorization': 'Bearer sb_publishable_VhddzdHirUzYfgMe7shYDA_-i7T4xnO'
}
req = urllib.request.Request('https://uexqmrspulhauubzrwyn.supabase.co/rest/v1/catalogo_insumos?select=*', headers=headers)
with urllib.request.urlopen(req) as resp:
    catalog = json.loads(resp.read().decode('utf-8'))

cat_by_id = {c['id']: c for c in catalog}
cat_by_name = {c['nombre'].lower().strip(): c for c in catalog}
print(f"Catálogo cargado: {len(catalog)} insumos.")

def get_insumo_info(name_or_id):
    if not name_or_id:
        return {'nombre': 'Desconocido', 'categoria': 'GENERAL', 'costo_unitario_kg': 0, 'peso_estandar_porcion_kg': 0.25}
    if name_or_id in cat_by_id:
        return cat_by_id[name_or_id]
    n = str(name_or_id).lower().strip()
    if n in cat_by_name:
        return cat_by_name[n]
    # Fuzzy / aliases
    aliases = {
        'damas': 'filete de pollo x 120',
        'filete pollo': 'filete de pollox 320 gr',
        'costilla res': 'costilla de res',
        'costilla blanca': 'costilla de res',
        'cordobue': 'cordon blue',
        'cordon blue': 'cordon blue',
        'francesa': 'papas a la francesa',
        'mix verdura': 'mix de verduras',
        'ahumada': 'costilla ahumada',
        'lomo cerdo': 'lomo de cerdo',
        'san luis': 'costilla san luis',
        'pernil pollo': 'pernil de pollo',
        'pernil cerdo': 'pernil de cerdo',
        'chorizo': 'chorizo paisa',
        'chorizos': 'chorizo paisa',
        'morcilla': 'rellena',
        'alguacil': 'pescado bagre',
        'bagre': 'pescado bagre',
        'salmon': 'salmon',
        'salmón': 'salmon',
        'pezuña': 'pezuña de cerdo',
        'caderita': 'caderita',
        'callo': 'callo',
        'sobrebarriga': 'sobrebarriga delgada',
        'huevos': 'huevo',
        'huevos aaa': 'huevo'
    }
    if n in aliases:
        target = aliases[n]
        if target in cat_by_name:
            return cat_by_name[target]
    # Substring search
    for k, v in cat_by_name.items():
        if n in k or k in n:
            return v
    return {'nombre': str(name_or_id).title(), 'categoria': 'CARNES', 'costo_unitario_kg': 20000, 'peso_estandar_porcion_kg': 0.25}

# -------------------------------------------------------------------------
# 2. RECOLECTAR TODOS LOS MOVIMIENTOS POR DÍA (01 AL 24 DE SEP)
# -------------------------------------------------------------------------
movements_by_day = collections.defaultdict(list)

# --- DÍA 01: En Supabase actualmente (14 operacionales + 9 saldos iniciales) ---
url_sep1 = 'https://uexqmrspulhauubzrwyn.supabase.co/rest/v1/movimientos_inventario?fecha=eq.2026-09-01&order=fecha_hora.asc'
req_sep1 = urllib.request.Request(url_sep1, headers=headers)
with urllib.request.urlopen(req_sep1) as resp:
    sep1_db = json.loads(resp.read().decode('utf-8'))

for m in sep1_db:
    c = cat_by_id.get(m.get('insumo_id'), {})
    movements_by_day['2026-09-01'].append({
        'id': m.get('id', ''),
        'fecha': '2026-09-01',
        'hora': str(m.get('fecha_hora', ''))[11:16] or '08:00',
        'tipo': m.get('tipo_movimiento'),
        'insumo': c.get('nombre', 'Insumo Sep 01'),
        'categoria': c.get('categoria', 'CARNE'),
        'origen': m.get('origen', 'BODEGA_ENTERO'),
        'destino': m.get('destino', 'COCINA_ENTERO'),
        'cant_kg': float(m.get('cant_sin_porcionar_kg') or 0),
        'porc_und': int(m.get('porciones_und') or 0),
        'peso_porc_kg': float(m.get('peso_porciones_kg') or 0),
        'costo_unit': float(m.get('costo_unitario_kg') or c.get('costo_unitario_kg') or 0),
        'costo_total': float(m.get('valor_total_movimiento') or 0),
        'observaciones': m.get('observaciones', ''),
        'estado': 'APROBADO',
        'fuente': 'Base de Datos Supabase (Línea Base Sep 01)'
    })

# --- DÍA 02: De insert_sep2_movs.js (19 movimientos) ---
with open('.archive/scratch_files/scratch/insert_sep2_movs.js', 'r', encoding='utf-8') as f:
    js_text = f.read()

# Extraer filasPlanilla
sep2_entries = re.findall(r"\{\s*tipo:\s*'([^']+)',\s*carne:\s*'([^']+)',\s*tipoEntrega:\s*'([^']+)',\s*cantidad:\s*([\d\.]+),\s*pesoKg:\s*([\d\.]+),\s*obs:\s*'([^']+)'\s*\}", js_text)
# Y también las compras y devoluciones que usan sintaxis extendida
sep2_all_blocks = re.findall(r"\{\s*tipo:\s*'([^']+)',\s*carne:\s*'([^']+)'([^}]+)\}", js_text)

for t, carne, rest in sep2_all_blocks:
    info = get_insumo_info(carne)
    c_kg_m = re.search(r"cantidadKg:\s*([\d\.]+)", rest)
    c_pkg_m = re.search(r"pesoKg:\s*([\d\.]+)", rest)
    cant_m = re.search(r"cantidad:\s*([\d\.]+)", rest)
    tot_m = re.search(r"costoTotal:\s*([\d\.]+)", rest)
    obs_m = re.search(r"obs:\s*'([^']+)'", rest)
    tipo_e_m = re.search(r"tipoEntrega:\s*'([^']+)'", rest)
    tipo_d_m = re.search(r"tipoDevolucion:\s*'([^']+)'", rest)
    
    cant_kg = float(c_kg_m.group(1)) if c_kg_m else 0.0
    peso_kg = float(c_pkg_m.group(1)) if c_pkg_m else 0.0
    cant_gen = float(cant_m.group(1)) if cant_m else 0.0
    obs = obs_m.group(1) if obs_m else ''
    
    is_porc = (tipo_e_m and tipo_e_m.group(1) == 'PORCIONADO') or (tipo_d_m and tipo_d_m.group(1) == 'PORCIONADO')
    
    porc_und = int(cant_gen) if is_porc else 0
    if not is_porc and cant_kg == 0.0 and cant_gen > 0:
        cant_kg = cant_gen
    if is_porc and peso_kg == 0.0 and porc_und > 0:
        peso_kg = porc_und * float(info.get('peso_estandar_porcion_kg') or 0.25)
        
    unit_cost = float(info.get('costo_unitario_kg') or 20000)
    costo_tot = float(tot_m.group(1)) if tot_m else ((cant_kg if cant_kg > 0 else peso_kg) * unit_cost)
    
    movements_by_day['2026-09-02'].append({
        'id': f"sep2-{len(movements_by_day['2026-09-02'])+1:02d}",
        'fecha': '2026-09-02',
        'hora': f"08:{len(movements_by_day['2026-09-02'])*3:02d}",
        'tipo': t,
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': 'PROVEEDOR' if 'COMPRA' in t else ('COCINA_PORCIONADO' if 'DEVOLUCION' in t else ('BODEGA_PORCIONADO' if is_porc else 'BODEGA_ENTERO')),
        'destino': 'BODEGA_ENTERO' if 'COMPRA' in t else ('BODEGA_PORCIONADO' if 'DEVOLUCION' in t else ('COCINA_PORCIONADO' if is_porc else 'COCINA_ENTERO')),
        'cant_kg': cant_kg,
        'porc_und': porc_und,
        'peso_porc_kg': peso_kg,
        'costo_unit': unit_cost,
        'costo_total': costo_tot,
        'observaciones': obs,
        'estado': 'PENDIENTE (Bodeguero)',
        'fuente': 'insert_sep2_movs.js'
    })

# --- DÍAS 03 Y 04: Compras de Excel + Salidas de Conciliación Operativa ---
# DÍA 03: 9 compras facturadas + salidas a cocina registradas
compras_d3 = [
    ('Caderita', 5.96, 34000, 202640, 'FE01576', 'Carnes Poblado'),
    ('Callo', 7.00, 20000, 140000, 'FE01576', 'Carnes Poblado'),
    ('Higado', 5.80, 22000, 127600, 'FE01570', 'Carnes Poblado'),
    ('Lomo viche', 7.35, 54000, 396900, 'FE01570', 'Carnes Poblado'),
    ('Caderita', 4.05, 33000, 133650, 'FE01569', 'Carnes Poblado'),
    ('Costilla Especial', 10.30, 20000, 206000, 'FE01569', 'Carnes Poblado'),
    ('Chorizo paisa', 4.00, 48000, 192000, '0000006666', 'Carnes Don Arturo'),
    ('Chorizo paisa', 1.00, 46000, 46000, '0000006666', 'Carnes Don Arturo'),
    ('Costilla de cerdo especial', 20.00, 18200, 364000, 'FE1-49657', 'Supertiendas Cañaveral')
]
for item, cant, unit, tot, fac, prov in compras_d3:
    info = get_insumo_info(item)
    movements_by_day['2026-09-03'].append({
        'id': f"sep3-in-{len(movements_by_day['2026-09-03'])+1:02d}",
        'fecha': '2026-09-03',
        'hora': '08:30',
        'tipo': 'ENTRADA_COMPRA',
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': f"PROVEEDOR ({prov})",
        'destino': 'BODEGA_ENTERO',
        'cant_kg': cant,
        'porc_und': 0,
        'peso_porc_kg': 0.0,
        'costo_unit': unit,
        'costo_total': tot,
        'observaciones': f"Entrada Factura {fac} - Proveedor: {prov}",
        'estado': 'APROBADO (Facturado)',
        'fuente': 'COSTOS_GASTOS_SEPTIEMBRE_2026.xlsx'
    })

# Salidas a cocina Día 03 (Reconciliación y planilla diaria)
salidas_d3 = [
    ('Baby beef', 'TRASLADO_COCINA', 'PORCIONADO', 0, 30, 10.5, 45000, 472500, 'Salida a cocina 30 porciones Baby Beef'),
    ('Churrasco', 'TRASLADO_COCINA', 'PORCIONADO', 0, 20, 7.0, 32000, 224000, 'Salida a cocina 20 porciones Churrasco'),
    ('Filete de pollo x 120', 'TRASLADO_COCINA', 'PORCIONADO', 0, 10, 2.5, 25000, 62500, 'Salida a cocina 10 porciones Filete Damas'),
    ('Callo', 'TRASLADO_COCINA', 'ENTERO', 7.0, 0, 0.0, 20000, 140000, 'Salida a cocina 7.0 Kg Callo para cocción'),
    ('Higado', 'TRASLADO_COCINA', 'ENTERO', 5.8, 0, 0.0, 22000, 127600, 'Salida a cocina 5.8 Kg Higado'),
    ('Lomo viche', 'TRASLADO_COCINA', 'ENTERO', 7.35, 0, 0.0, 54000, 396900, 'Salida a cocina 7.35 Kg Lomo Viche'),
    ('Costilla de cerdo especial', 'TRASLADO_COCINA', 'ENTERO', 15.0, 0, 0.0, 18200, 273000, 'Salida a cocina 15.0 Kg Costilla Cerdo'),
    ('Chorizo paisa', 'TRASLADO_COCINA', 'PORCIONADO', 0, 10, 2.5, 48000, 120000, 'Salida a cocina 10 porciones Chorizo'),
    ('Higado', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 24, 7.99, 22000, 175780, 'Devolución de cocina 24 porciones Higado porcionado'),
    ('Lomo viche', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 15, 5.25, 54000, 283500, 'Devolución de cocina 15 porciones Lomo Viche porcionado'),
    ('Costilla Especial', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 20, 5.72, 20000, 114400, 'Devolución de cocina 20 porciones Costilla Especial'),
    ('Pernil de pollo', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 40, 13.3, 10000, 133000, 'Devolución de cocina 40 porciones Pernil Pollo')
]
for item, t, modo, ckg, pund, ppkg, unit, tot, obs in salidas_d3:
    info = get_insumo_info(item)
    movements_by_day['2026-09-03'].append({
        'id': f"sep3-op-{len(movements_by_day['2026-09-03'])+1:02d}",
        'fecha': '2026-09-03',
        'hora': '10:00',
        'tipo': t,
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': 'COCINA_PORCIONADO' if 'DEVOLUCION' in t else ('BODEGA_PORCIONADO' if modo == 'PORCIONADO' else 'BODEGA_ENTERO'),
        'destino': 'BODEGA_PORCIONADO' if 'DEVOLUCION' in t else ('COCINA_PORCIONADO' if modo == 'PORCIONADO' else 'COCINA_ENTERO'),
        'cant_kg': ckg,
        'porc_und': pund,
        'peso_porc_kg': ppkg,
        'costo_unit': unit,
        'costo_total': tot,
        'observaciones': obs,
        'estado': 'PENDIENTE (Bodeguero)',
        'fuente': 'Planilla Operativa / Reconciliación Maestra Sep 03'
    })

# DÍA 04: 8 compras facturadas + movimientos operativos y de pescado
compras_d4 = [
    ('Lomo de cerdo', 26.50, 19500, 516750, 'FEBF-33', 'Supertiendas Cañaveral'),
    ('Filete de pollo x 350', 23.50, 18500, 434750, 'FAC-920', 'Alexis Valencia'),
    ('Rellena', 10.00, 10000, 100000, 'FAC-8118', 'Carnes y Embutidos'),
    ('Pechuga mcpollo', 3.095, 13980, 43269, '05FZ91332', 'Supertiendas Cañaveral'),
    ('Tilapia 350/450', 20.74, 18000, 373320, 'FVE 130', 'Frio Sabor / Gabriel Ramirez'),
    ('Salmon', 6.765, 55000, 372075, 'FVE 130', 'Frio Sabor / Gabriel Ramirez'),
    ('Trucha', 9.16, 23000, 210680, 'FVE 130', 'Frio Sabor / Gabriel Ramirez'),
    ('Marisco', 7.00, 24000, 168000, 'FVE 130', 'Frio Sabor / Gabriel Ramirez')
]
for item, cant, unit, tot, fac, prov in compras_d4:
    info = get_insumo_info(item)
    movements_by_day['2026-09-04'].append({
        'id': f"sep4-in-{len(movements_by_day['2026-09-04'])+1:02d}",
        'fecha': '2026-09-04',
        'hora': '08:45',
        'tipo': 'ENTRADA_COMPRA',
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': f"PROVEEDOR ({prov})",
        'destino': 'BODEGA_ENTERO',
        'cant_kg': cant,
        'porc_und': 0,
        'peso_porc_kg': 0.0,
        'costo_unit': unit,
        'costo_total': tot,
        'observaciones': f"Entrada Factura {fac} - Proveedor: {prov}",
        'estado': 'APROBADO (Facturado)',
        'fuente': 'COSTOS_GASTOS_SEPTIEMBRE_2026.xlsx'
    })

# Salidas y devoluciones Día 04
salidas_d4 = [
    ('Salmon', 'ENTRADA_COMPRA', 'ENTERO', 6.7, 0, 0.0, 55000, 368500, '[PENDIENTE_APROBAR] Entrada del día 6.7 Kg Salmon (Planilla 04-Sep)'),
    ('Salmon', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 21, 4.2, 55000, 231000, '[PENDIENTE_APROBAR] Devolución de cocina 21 porciones Salmon (Planilla 04-Sep)'),
    ('Tilapia 350/450', 'TRASLADO_COCINA', 'PORCIONADO', 0, 20, 8.0, 18000, 144000, 'Salida a cocina 20 porciones Tilapia'),
    ('Trucha', 'TRASLADO_COCINA', 'PORCIONADO', 0, 15, 4.5, 23000, 103500, 'Salida a cocina 15 porciones Trucha'),
    ('Filete de pollox 320 gr', 'TRASLADO_COCINA', 'PORCIONADO', 0, 35, 11.2, 18500, 207200, 'Salida a cocina 35 porciones Filete Pollo 320g'),
    ('Lomo de cerdo', 'TRASLADO_COCINA', 'ENTERO', 20.0, 0, 0.0, 19500, 390000, 'Salida a cocina 20.0 Kg Lomo de Cerdo'),
    ('Lomo de cerdo', 'DEVOLUCION_COCINA', 'PORCIONADO', 0, 65, 14.43, 19500, 281385, 'Devolución de cocina 65 porciones Lomo Cerdo porcionado'),
    ('Marisco', 'TRASLADO_COCINA', 'ENTERO', 5.0, 0, 0.0, 24000, 120000, 'Salida a cocina 5.0 Kg Mariscos'),
    ('Rellena', 'TRASLADO_COCINA', 'ENTERO', 10.0, 0, 0.0, 10000, 100000, 'Salida a cocina 10.0 Kg Rellena')
]
for item, t, modo, ckg, pund, ppkg, unit, tot, obs in salidas_d4:
    info = get_insumo_info(item)
    movements_by_day['2026-09-04'].append({
        'id': f"sep4-op-{len(movements_by_day['2026-09-04'])+1:02d}",
        'fecha': '2026-09-04',
        'hora': '10:15',
        'tipo': t,
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': 'COCINA_PORCIONADO' if 'DEVOLUCION' in t else ('BODEGA_PORCIONADO' if modo == 'PORCIONADO' else 'BODEGA_ENTERO'),
        'destino': 'BODEGA_PORCIONADO' if 'DEVOLUCION' in t else ('COCINA_PORCIONADO' if modo == 'PORCIONADO' else 'COCINA_ENTERO'),
        'cant_kg': ckg,
        'porc_und': pund,
        'peso_porc_kg': ppkg,
        'costo_unit': unit,
        'costo_total': tot,
        'observaciones': obs,
        'estado': 'PENDIENTE (Bodeguero)',
        'fuente': 'Planilla Operativa / Step 601 Sep 04'
    })

# --- DÍAS 05 AL 22: Desde los scripts de seed estructurados ---
seed_files = [
    'scripts/seed_movimientos_bodeguero.mjs',
    'scripts/seed_movimientos_07_10_11_sep.mjs',
    'scripts/seed_movimientos_09_sep.mjs',
    'scripts/seed_movimientos_12_13_15_16_sep.mjs',
    'scripts/seed_movimientos_20_21_22_sep.mjs',
]

for sf in seed_files:
    with open(sf, 'r', encoding='utf-8') as f:
        content = f.read()
    
    matches = re.finditer(r"\{\s*fecha:\s*'([^']+)',\s*insumo:\s*'([^']+)',\s*tipo:\s*'([^']+)'([^}]+)\}", content)
    for m in matches:
        d = m.group(1)
        insumo = m.group(2)
        tipo = m.group(3)
        rest = m.group(4)
        
        cant_kg_m = re.search(r"cantKg:\s*([\d\.]+)", rest)
        porc_m = re.search(r"porc:\s*(\d+)", rest)
        obs_m = re.search(r"obs:\s*'([^']+)'", rest)
        
        cant_kg = float(cant_kg_m.group(1)) if cant_kg_m else 0.0
        porc = int(porc_m.group(1)) if porc_m else 0
        obs = obs_m.group(1) if obs_m else ''
        
        info = get_insumo_info(insumo)
        unit_cost = float(info.get('costo_unitario_kg') or 20000)
        std_weight = float(info.get('peso_estandar_porcion_kg') or 0.25)
        peso_porc = porc * std_weight if porc > 0 else 0.0
        tot_cost = (cant_kg if cant_kg > 0 else peso_porc) * unit_cost
        
        movements_by_day[d].append({
            'id': f"{d}-{len(movements_by_day[d])+1:02d}",
            'fecha': d,
            'hora': f"09:{len(movements_by_day[d])*2 % 60:02d}",
            'tipo': tipo,
            'insumo': info['nombre'],
            'categoria': info['categoria'],
            'origen': 'PROVEEDOR' if 'COMPRA' in tipo else ('COCINA_PORCIONADO' if 'DEVOLUCION' in tipo else ('BODEGA_PORCIONADO' if porc > 0 else 'BODEGA_ENTERO')),
            'destino': 'BODEGA_ENTERO' if 'COMPRA' in tipo else ('BODEGA_PORCIONADO' if 'DEVOLUCION' in tipo else ('COCINA_PORCIONADO' if porc > 0 else 'COCINA_ENTERO')),
            'cant_kg': cant_kg,
            'porc_und': porc,
            'peso_porc_kg': peso_porc,
            'costo_unit': unit_cost,
            'costo_total': tot_cost,
            'observaciones': obs,
            'estado': 'PENDIENTE (Bodeguero)',
            'fuente': os.path.basename(sf)
        })

# --- DÍA 23: Entrada Facturada Tocineta Barril + Porcionados Operativos ---
info_tocineta = get_insumo_info('tocineta barril')
movements_by_day['2026-09-23'].append({
    'id': 'sep23-01',
    'fecha': '2026-09-23',
    'hora': '08:30',
    'tipo': 'ENTRADA_COMPRA',
    'insumo': info_tocineta['nombre'],
    'categoria': info_tocineta['categoria'],
    'origen': 'PROVEEDOR (Carnes Los Monos)',
    'destino': 'BODEGA_ENTERO',
    'cant_kg': 64.8,
    'porc_und': 0,
    'peso_porc_kg': 0.0,
    'costo_unit': 21900.0,
    'costo_total': 1419120.0,
    'observaciones': 'Entrada Factura YP24600 Los Monos - 64.8 kg tocineta barril ("entró ayer" anotado en conteo del 24)',
    'estado': 'APROBADO',
    'fuente': 'apply_full_master_integration.py'
})

# Porcionados maestros de conciliación operativa aplicados en septiembre
porcionados_maestros = [
    ('2026-09-01', 'Morrillo', 14.0, 60, 12.0, 'PORCIONADO', 'Porcionado 14 kg a 60 porciones de morrillo (2 kg merma)'),
    ('2026-09-02', 'Cordon Blue', 0.0, 40, 11.4, 'AJUSTE_ENTRADA', 'Porcionado de 40 porciones cordon blue desde pechuga'),
    ('2026-09-05', 'Punta de anca', 39.2, 112, 39.2, 'PORCIONADO', 'Porcionado operativo lote Punta de Anca a 112 porciones'),
    ('2026-09-05', 'Churrasco', 38.5, 110, 38.5, 'PORCIONADO', 'Porcionado operativo lote Churrasco a 110 porciones'),
    ('2026-09-05', 'Baby beef', 14.35, 41, 14.35, 'PORCIONADO', 'Porcionado operativo lote Baby Beef a 41 porciones'),
    ('2026-09-05', 'Costilla san luis', 52.15, 149, 52.15, 'PORCIONADO', 'Porcionado Costilla San Luis a 149 porciones'),
    ('2026-09-06', 'Gallina', 86.7, 222, 86.7, 'PORCIONADO', 'Porcionado de Gallina a 222 presas/cuartos para sancocho'),
    ('2026-09-07', 'Pescado bagre', 25.9, 74, 25.9, 'PORCIONADO', 'Porcionado de Bagre a 74 porciones'),
    ('2026-09-07', 'Trucha', 22.2, 74, 22.2, 'PORCIONADO', 'Porcionado de Trucha a 74 porciones')
]
for d, carne, ckg, pund, ppkg, t, obs in porcionados_maestros:
    info = get_insumo_info(carne)
    unit_cost = float(info.get('costo_unitario_kg') or 20000)
    movements_by_day[d].append({
        'id': f"porc-{d}-{info['codigo'] if 'codigo' in info else info['nombre'][:3]}",
        'fecha': d,
        'hora': '14:00',
        'tipo': t,
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': 'BODEGA_ENTERO',
        'destino': 'BODEGA_PORCIONADO',
        'cant_kg': ckg,
        'porc_und': pund,
        'peso_porc_kg': ppkg,
        'costo_unit': unit_cost,
        'costo_total': (ckg if ckg > 0 else ppkg) * unit_cost,
        'observaciones': obs,
        'estado': 'APROBADO (Ajuste Operativo)',
        'fuente': 'simulate_full_reconciliation.py'
    })

# --- DÍA 24: Movimientos de Bodeguero del 24 + Ajustes de Auditoría Física ---
# 1. Movimientos operacionales de bodega (18 movimientos)
import sys
sys.path.append('scratch')
try:
    import generate_sql
    for m in generate_sql.movimientos:
        if m['fecha'] == '2026-09-24':
            c = cat_by_id.get(m['insumo_id'], {})
            movements_by_day['2026-09-24'].append({
                'id': f"sep24-bod-{len(movements_by_day['2026-09-24'])+1:02d}",
                'fecha': '2026-09-24',
                'hora': str(m.get('fecha_hora', ''))[11:16] or '10:00',
                'tipo': m['tipo_movimiento'],
                'insumo': c.get('nombre', 'Insumo 24-Sep'),
                'categoria': c.get('categoria', 'CARNE'),
                'origen': m['origen'],
                'destino': m['destino'],
                'cant_kg': float(m['cant_sin_porcionar_kg']),
                'porc_und': int(m['porciones_und']),
                'peso_porc_kg': float(m['peso_porciones_kg']),
                'costo_unit': float(m['costo_unitario_kg']),
                'costo_total': float(m['valor_total_movimiento']),
                'observaciones': m['observaciones'],
                'estado': 'PENDIENTE (Bodeguero)',
                'fuente': 'scratch/generate_sql.py'
            })
except Exception as e:
    print(f"Nota con generate_sql: {e}")

# 2. Ajustes de Inconsistencias de Auditoría Física del 24-Sep (20 ajustes auditados)
ajustes_24_sep = [
    ('Morrillo', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -10, -2.5, 28000, -70000, 'Ajuste conteo físico 24-Sep (Diferencia -10 porciones)'),
    ('Costilla de res', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -15, -4.5, 20000, -90000, 'Ajuste conteo físico 24-Sep (Diferencia -15 porciones)'),
    ('Sobrebarriga delgada', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -8, -2.4, 25000, -60000, 'Ajuste conteo físico 24-Sep (Diferencia -8 porciones)'),
    ('Lengua de res', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -5, -1.5, 23900, -35850, 'Ajuste conteo físico 24-Sep (Diferencia -5 porciones)'),
    ('Punta de anca', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -7, -2.45, 33000, -80850, 'Ajuste conteo físico 24-Sep (Diferencia -7 porciones)'),
    ('Churrasco', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -6, -2.1, 32000, -67200, 'Ajuste conteo físico 24-Sep (Diferencia -6 porciones)'),
    ('Baby beef', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -4, -1.4, 45000, -63000, 'Ajuste conteo físico 24-Sep (Diferencia -4 porciones)'),
    ('Costilla san luis', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -9, -3.15, 21900, -68985, 'Ajuste conteo físico 24-Sep (Diferencia -9 porciones)'),
    ('Lomo de cerdo', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -12, -2.66, 19500, -51870, 'Ajuste conteo físico 24-Sep (Diferencia -12 porciones)'),
    ('Pernil de pollo', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -14, -4.66, 10000, -46600, 'Ajuste conteo físico 24-Sep (Diferencia -14 porciones)'),
    ('Filete de pollox 320 gr', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -11, -3.52, 18500, -65120, 'Ajuste conteo físico 24-Sep (Diferencia -11 porciones)'),
    ('Gallina', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -2, -0.78, 51933, -40507, 'Ajuste conteo físico 24-Sep (Diferencia -2 unidades)'),
    ('Pescado bagre', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -6, -2.1, 14000, -29400, 'Ajuste conteo físico 24-Sep (Diferencia -6 porciones)'),
    ('Trucha', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -5, -1.5, 23000, -34500, 'Ajuste conteo físico 24-Sep (Diferencia -5 porciones)'),
    ('Salmon', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -4, -0.8, 55000, -44000, 'Ajuste conteo físico 24-Sep (Diferencia -4 porciones)'),
    ('Tilapia 350/450', 'AJUSTE_INVENTARIO', 'BODEGA_PORCIONADO', 'COCINA_PORCIONADO', 0.0, -8, -3.2, 18000, -57600, 'Ajuste conteo físico 24-Sep (Diferencia -8 porciones)'),
    ('Costilla ahumada', 'AJUSTE_INVENTARIO', 'BODEGA_ENTERO', 'COCINA_ENTERO', -3.2, 0, 0.0, 24000, -76800, 'Ajuste conteo físico 24-Sep (Diferencia -3.2 Kg)'),
    ('Tocino', 'AJUSTE_INVENTARIO', 'BODEGA_ENTERO', 'COCINA_ENTERO', -5.4, 0, 0.0, 21900, -118260, 'Ajuste conteo físico 24-Sep (Diferencia -5.4 Kg)'),
    ('Sobaco', 'AJUSTE_INVENTARIO', 'BODEGA_ENTERO', 'COCINA_ENTERO', -4.1, 0, 0.0, 25900, -106190, 'Ajuste conteo físico 24-Sep (Diferencia -4.1 Kg)'),
    ('Tocineta barril', 'AJUSTE_INVENTARIO', 'BODEGA_ENTERO', 'COCINA_ENTERO', -6.0, 0, 0.0, 21900, -131400, 'Ajuste conteo físico 24-Sep (Diferencia -6.0 Kg)')
]
for carne, t, org, dst, ckg, pund, ppkg, unit, tot, obs in ajustes_24_sep:
    info = get_insumo_info(carne)
    movements_by_day['2026-09-24'].append({
        'id': f"sep24-ajuste-{len(movements_by_day['2026-09-24'])+1:02d}",
        'fecha': '2026-09-24',
        'hora': '11:00',
        'tipo': t,
        'insumo': info['nombre'],
        'categoria': info['categoria'],
        'origen': org,
        'destino': dst,
        'cant_kg': ckg,
        'porc_und': pund,
        'peso_porc_kg': ppkg,
        'costo_unit': unit,
        'costo_total': tot,
        'observaciones': obs,
        'estado': 'AUDITORÍA CONTEO FÍSICO',
        'fuente': 'Auditoría Conteo 24-Sep (Eliminado)'
    })

print("\n--- RESUMEN FINAL DE MOVIMIENTOS A ESCRIBIR EN EXCEL ---")
total_mes = 0
for d in sorted(movements_by_day.keys()):
    c = len(movements_by_day[d])
    total_mes += c
    print(f"  {d}: {c:3d} movimientos")
print(f"TOTAL GENERAL: {total_mes} movimientos en 24 días.")

# -------------------------------------------------------------------------
# 3. CREAR WORKBOOK DE OPENPYXL CON DISEÑO PROFESIONAL
# -------------------------------------------------------------------------
wb = openpyxl.Workbook()
wb.remove(wb.active) # Quitar hoja por defecto

# Paleta de colores corporativa
C_HEADER_BG = "1B365D"    # Azul Marino Corporativo
C_SUB_BG = "2C4D75"       # Azul Medio
C_SECTION_BG = "3E638D"   # Azul Claro
C_TOTAL_BG = "D9E2EC"     # Gris Azulado
C_ZEBRA = "F8F9FA"        # Gris tenue
C_BORDER = "C4CBD5"

font_title = Font(name="Calibri", size=14, bold=True, color="1B365D")
font_sub = Font(name="Calibri", size=9, italic=True, color="4A5568")
font_kpi_val = Font(name="Calibri", size=11, bold=True, color="1B365D")
font_col = Font(name="Calibri", size=8.5, bold=True, color="FFFFFF")
font_cell = Font(name="Calibri", size=8, color="111111")
font_cell_bold = Font(name="Calibri", size=8, bold=True, color="111111")
font_total = Font(name="Calibri", size=8.5, bold=True, color="1B365D")

fill_hdr = PatternFill("solid", fgColor=C_HEADER_BG)
fill_sub = PatternFill("solid", fgColor=C_SUB_BG)
fill_total = PatternFill("solid", fgColor=C_TOTAL_BG)
fill_zebra = PatternFill("solid", fgColor=C_ZEBRA)

thin_border = Border(
    left=Side(style='thin', color=C_BORDER),
    right=Side(style='thin', color=C_BORDER),
    top=Side(style='thin', color=C_BORDER),
    bottom=Side(style='thin', color=C_BORDER)
)
total_border = Border(
    left=Side(style='thin', color=C_BORDER),
    right=Side(style='thin', color=C_BORDER),
    top=Side(style='thin', color=C_BORDER),
    bottom=Side(style='double', color="1B365D")
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# =========================================================================
# HOJA 1: RESUMEN_EJECUTIVO_MES
# =========================================================================
ws_resumen = wb.create_sheet(title="RESUMEN_EJECUTIVO")
ws_resumen.views.sheetView[0].showGridLines = True

ws_resumen.merge_cells("A1:M1")
ws_resumen["A1"] = "RESTAURANTE CAMPESTRE LA FINCA — HISTORIAL CONSOLIDADO DE MOVIMIENTOS DE BODEGA"
ws_resumen["A1"].font = font_title
ws_resumen["A1"].alignment = align_center
ws_resumen.row_dimensions[1].height = 22

ws_resumen.merge_cells("A2:M2")
ws_resumen["A2"] = "Auditoría Histórica Completa: Movimientos Guardados y Auditados entre el 01 y el 24 de Septiembre de 2026"
ws_resumen["A2"].font = font_sub
ws_resumen["A2"].alignment = align_center
ws_resumen.row_dimensions[2].height = 15

headers_resumen = [
    ("A4", "Fecha"),
    ("B4", "Total Movs"),
    ("C4", "N° Compras"),
    ("D4", "Kg Comprados"),
    ("E4", "Inversión Compras ($)"),
    ("F4", "N° Traslados Cocina"),
    ("G4", "Kg Trasladados"),
    ("H4", "Porciones Cocina (Und)"),
    ("I4", "Costo Traslados ($)"),
    ("J4", "N° Devoluciones"),
    ("K4", "Porciones Devueltas (Und)"),
    ("L4", "N° Porcionados / Ajustes"),
    ("M4", "Estado de Jornada en Sistema")
]
ws_resumen.row_dimensions[4].height = 20
for pos, txt in headers_resumen:
    c = ws_resumen[pos]
    c.value = txt
    c.font = font_col
    c.fill = fill_hdr
    c.alignment = align_center
    c.border = thin_border

row_idx = 5
for d in sorted(movements_by_day.keys()):
    movs = movements_by_day[d]
    compras = [m for m in movs if 'COMPRA' in m['tipo']]
    traslados = [m for m in movs if 'TRASLADO' in m['tipo']]
    devoluciones = [m for m in movs if 'DEVOLUCION' in m['tipo']]
    otros = [m for m in movs if 'COMPRA' not in m['tipo'] and 'TRASLADO' not in m['tipo'] and 'DEVOLUCION' not in m['tipo']]
    
    kg_comp = sum(m['cant_kg'] for m in compras)
    cost_comp = sum(m['costo_total'] for m in compras)
    kg_tras = sum(m['cant_kg'] for m in traslados)
    und_tras = sum(m['porc_und'] for m in traslados)
    cost_tras = sum(m['costo_total'] for m in traslados)
    und_dev = sum(m['porc_und'] for m in devoluciones)
    
    ws_resumen.row_dimensions[row_idx].height = 15
    is_even = (row_idx % 2 == 0)
    rf = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
    
    vals = [
        (d, align_center, "@", font_cell_bold),
        (len(movs), align_center, "#,##0", font_cell_bold),
        (len(compras), align_center, "#,##0", font_cell),
        (kg_comp, align_right, "#,##0.00 \"kg\"", font_cell),
        (cost_comp, align_right, "$#,##0", font_cell),
        (len(traslados), align_center, "#,##0", font_cell),
        (kg_tras, align_right, "#,##0.00 \"kg\"", font_cell),
        (und_tras, align_center, "#,##0 \"und\"", font_cell),
        (cost_tras, align_right, "$#,##0", font_cell),
        (len(devoluciones), align_center, "#,##0", font_cell),
        (und_dev, align_center, "#,##0 \"und\"", font_cell),
        (len(otros), align_center, "#,##0", font_cell),
        ("APROBADO" if d == '2026-09-01' else ("AUDITORÍA / AJUSTE" if d == '2026-09-24' else "DEPURADO / RESPALDADO"), align_center, "@", font_cell_bold)
    ]
    
    for col_i, (v, aln, fmt, fnt) in enumerate(vals, start=1):
        cell = ws_resumen.cell(row=row_idx, column=col_i, value=v)
        cell.alignment = aln
        cell.number_format = fmt
        cell.font = fnt
        cell.fill = rf
        cell.border = thin_border
        
    row_idx += 1

# Fila de Totales
ws_resumen.row_dimensions[row_idx].height = 18
tot_cols = [
    ("TOTALES MES", align_center, "@"),
    (f"=SUM(B5:B{row_idx-1})", align_center, "#,##0"),
    (f"=SUM(C5:C{row_idx-1})", align_center, "#,##0"),
    (f"=SUM(D5:D{row_idx-1})", align_right, "#,##0.00 \"kg\""),
    (f"=SUM(E5:E{row_idx-1})", align_right, "$#,##0"),
    (f"=SUM(F5:F{row_idx-1})", align_center, "#,##0"),
    (f"=SUM(G5:G{row_idx-1})", align_right, "#,##0.00 \"kg\""),
    (f"=SUM(H5:H{row_idx-1})", align_center, "#,##0 \"und\""),
    (f"=SUM(I5:I{row_idx-1})", align_right, "$#,##0"),
    (f"=SUM(J5:J{row_idx-1})", align_center, "#,##0"),
    (f"=SUM(K5:K{row_idx-1})", align_center, "#,##0 \"und\""),
    (f"=SUM(L5:L{row_idx-1})", align_center, "#,##0"),
    ("HISTORIAL 01-24 SEP", align_center, "@")
]
for col_i, (formula_or_txt, aln, fmt) in enumerate(tot_cols, start=1):
    cell = ws_resumen.cell(row=row_idx, column=col_i, value=formula_or_txt)
    cell.alignment = aln
    cell.number_format = fmt
    cell.font = font_total
    cell.fill = fill_total
    cell.border = total_border

# Anchos de columna resumen
col_widths_res = {
    'A': 13, 'B': 12, 'C': 12, 'D': 15, 'E': 18,
    'F': 16, 'G': 15, 'H': 18, 'I': 18, 'J': 15,
    'K': 19, 'L': 18, 'M': 24
}
for c, w in col_widths_res.items():
    ws_resumen.column_dimensions[c].width = w

# =========================================================================
# HOJAS DÍA POR DÍA (01_Sep AL 24_Sep)
# =========================================================================
headers_dia = [
    ("A3", "ID"),
    ("B3", "Fecha"),
    ("C3", "Hora"),
    ("D3", "Tipo de Movimiento"),
    ("E3", "Insumo / Corte"),
    ("F3", "Categoría"),
    ("G3", "Origen"),
    ("H3", "Destino"),
    ("I3", "Cant. Entero (Kg)"),
    ("J3", "Porciones (Und)"),
    ("K3", "Peso Porc. (Kg)"),
    ("L3", "Costo Unit. ($)"),
    ("M3", "Costo Total Mov. ($)"),
    ("N3", "Observaciones / Detalle"),
    ("O3", "Estado Operativo Previo"),
    ("P3", "Fuente Técnica")
]

col_widths_dia = {
    'A': 14, 'B': 12, 'C': 8, 'D': 20, 'E': 25,
    'F': 18, 'G': 20, 'H': 20, 'I': 16, 'J': 15,
    'K': 15, 'L': 15, 'M': 18, 'N': 42, 'O': 20, 'P': 24
}

for d in sorted(movements_by_day.keys()):
    day_num = d.split('-')[2]
    sheet_title = f"{day_num}_Sep"
    ws_d = wb.create_sheet(title=sheet_title)
    ws_d.views.sheetView[0].showGridLines = True
    
    # Encabezado
    ws_d.merge_cells("A1:P1")
    ws_d["A1"] = f"RESTAURANTE CAMPESTRE LA FINCA — MOVIMIENTOS REGISTRADOS DÍA {day_num} DE SEPTIEMBRE DE 2026"
    ws_d["A1"].font = font_title
    ws_d["A1"].alignment = align_center
    ws_d.row_dimensions[1].height = 20
    
    ws_d.merge_cells("A2:P2")
    ws_d["A2"] = f"Detalle de transacciones de bodega, traslados a cocina, devoluciones y porcionados para la fecha: {d}"
    ws_d["A2"].font = font_sub
    ws_d["A2"].alignment = align_center
    ws_d.row_dimensions[2].height = 14
    
    ws_d.row_dimensions[3].height = 20
    for pos, txt in headers_dia:
        c = ws_d[pos]
        c.value = txt
        c.font = font_col
        c.fill = fill_hdr
        c.alignment = align_center
        c.border = thin_border
        
    movs = movements_by_day[d]
    d_row = 4
    for m in movs:
        ws_d.row_dimensions[d_row].height = 14
        is_even = (d_row % 2 == 0)
        rf = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
        
        row_vals = [
            (m['id'][:12], align_center, "@", font_cell),
            (m['fecha'], align_center, "@", font_cell_bold),
            (m['hora'], align_center, "@", font_cell),
            (m['tipo'], align_center, "@", font_cell_bold),
            (m['insumo'], align_left, "@", font_cell_bold),
            (m['categoria'], align_left, "@", font_cell),
            (m['origen'], align_left, "@", font_cell),
            (m['destino'], align_left, "@", font_cell),
            (m['cant_kg'], align_right, "#,##0.00 \"kg\"", font_cell),
            (m['porc_und'], align_center, "#,##0 \"und\"", font_cell),
            (m['peso_porc_kg'], align_right, "#,##0.00 \"kg\"", font_cell),
            (m['costo_unit'], align_right, "$#,##0", font_cell),
            (m['costo_total'], align_right, "$#,##0", font_cell_bold),
            (m['observaciones'], align_left, "@", font_cell),
            (m['estado'], align_center, "@", font_cell_bold),
            (m['fuente'], align_left, "@", font_cell)
        ]
        
        for col_i, (v, aln, fmt, fnt) in enumerate(row_vals, start=1):
            cell = ws_d.cell(row=d_row, column=col_i, value=v)
            cell.alignment = aln
            cell.number_format = fmt
            cell.font = fnt
            cell.fill = rf
            cell.border = thin_border
            
        d_row += 1
        
    # Totales del día
    ws_d.row_dimensions[d_row].height = 18
    tot_d_vals = [
        ("TOTAL DÍA", align_center, "@"),
        ("", align_center, "@"),
        ("", align_center, "@"),
        (f"{len(movs)} MOVIMIENTOS", align_center, "@"),
        ("", align_center, "@"),
        ("", align_center, "@"),
        ("", align_center, "@"),
        ("", align_center, "@"),
        (f"=SUM(I4:I{d_row-1})", align_right, "#,##0.00 \"kg\""),
        (f"=SUM(J4:J{d_row-1})", align_center, "#,##0 \"und\""),
        (f"=SUM(K4:K{d_row-1})", align_right, "#,##0.00 \"kg\""),
        ("", align_center, "@"),
        (f"=SUM(M4:M{d_row-1})", align_right, "$#,##0"),
        ("", align_center, "@"),
        ("", align_center, "@"),
        ("", align_center, "@")
    ]
    for col_i, (v, aln, fmt) in enumerate(tot_d_vals, start=1):
        cell = ws_d.cell(row=d_row, column=col_i, value=v)
        cell.alignment = aln
        cell.number_format = fmt
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border
        
    for c, w in col_widths_dia.items():
        ws_d.column_dimensions[c].width = w

# =========================================================================
# HOJA FINAL: CONSOLIDADO_TODO_EL_MES (Tabla Continua)
# =========================================================================
ws_all = wb.create_sheet(title="CONSOLIDADO_TODO_EL_MES")
ws_all.views.sheetView[0].showGridLines = True

ws_all.merge_cells("A1:P1")
ws_all["A1"] = "RESTAURANTE CAMPESTRE LA FINCA — CONSOLIDADO MAESTRO DE TODOS LOS MOVIMIENTOS DE SEPTIEMBRE (01 AL 24)"
ws_all["A1"].font = font_title
ws_all["A1"].alignment = align_center
ws_all.row_dimensions[1].height = 20

ws_all.row_dimensions[2].height = 20
for pos, txt in [(p.replace('3', '2'), t) for p, t in headers_dia]:
    c = ws_all[pos]
    c.value = txt
    c.font = font_col
    c.fill = fill_hdr
    c.alignment = align_center
    c.border = thin_border

all_row = 3
for d in sorted(movements_by_day.keys()):
    for m in movements_by_day[d]:
        ws_all.row_dimensions[all_row].height = 14
        is_even = (all_row % 2 == 0)
        rf = fill_zebra if is_even else PatternFill("solid", fgColor="FFFFFF")
        
        row_vals = [
            (m['id'][:12], align_center, "@", font_cell),
            (m['fecha'], align_center, "@", font_cell_bold),
            (m['hora'], align_center, "@", font_cell),
            (m['tipo'], align_center, "@", font_cell_bold),
            (m['insumo'], align_left, "@", font_cell_bold),
            (m['categoria'], align_left, "@", font_cell),
            (m['origen'], align_left, "@", font_cell),
            (m['destino'], align_left, "@", font_cell),
            (m['cant_kg'], align_right, "#,##0.00 \"kg\"", font_cell),
            (m['porc_und'], align_center, "#,##0 \"und\"", font_cell),
            (m['peso_porc_kg'], align_right, "#,##0.00 \"kg\"", font_cell),
            (m['costo_unit'], align_right, "$#,##0", font_cell),
            (m['costo_total'], align_right, "$#,##0", font_cell_bold),
            (m['observaciones'], align_left, "@", font_cell),
            (m['estado'], align_center, "@", font_cell_bold),
            (m['fuente'], align_left, "@", font_cell)
        ]
        
        for col_i, (v, aln, fmt, fnt) in enumerate(row_vals, start=1):
            cell = ws_all.cell(row=all_row, column=col_i, value=v)
            cell.alignment = aln
            cell.number_format = fmt
            cell.font = fnt
            cell.fill = rf
            cell.border = thin_border
            
        all_row += 1

# Total Consolidado
ws_all.row_dimensions[all_row].height = 18
tot_all_vals = [
    ("TOTAL MES", align_center, "@"),
    ("", align_center, "@"),
    ("", align_center, "@"),
    (f"{total_mes} REGISTROS", align_center, "@"),
    ("", align_center, "@"),
    ("", align_center, "@"),
    ("", align_center, "@"),
    ("", align_center, "@"),
    (f"=SUM(I3:I{all_row-1})", align_right, "#,##0.00 \"kg\""),
    (f"=SUM(J3:J{all_row-1})", align_center, "#,##0 \"und\""),
    (f"=SUM(K3:K{all_row-1})", align_right, "#,##0.00 \"kg\""),
    ("", align_center, "@"),
    (f"=SUM(M3:M{all_row-1})", align_right, "$#,##0"),
    ("", align_center, "@"),
    ("", align_center, "@"),
    ("", align_center, "@")
]
for col_i, (v, aln, fmt) in enumerate(tot_all_vals, start=1):
    cell = ws_all.cell(row=all_row, column=col_i, value=v)
    cell.alignment = aln
    cell.number_format = fmt
    cell.font = font_total
    cell.fill = fill_total
    cell.border = total_border

for c, w in col_widths_dia.items():
    ws_all.column_dimensions[c].width = w

# -------------------------------------------------------------------------
# 4. GUARDAR ARCHIVO EXCEL
# -------------------------------------------------------------------------
output_path = "DOCUMENTOS BASE/MOVIMIENTOS_HISTORICOS_SEPTIEMBRE_2026.xlsx"
wb.save(output_path)
print(f"\n[OK] Archivo Excel generado exitosamente en: {output_path}")
print(f"Total pestanas creadas ({len(wb.sheetnames)}): {wb.sheetnames}")
