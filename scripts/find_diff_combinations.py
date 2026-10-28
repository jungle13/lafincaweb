import openpyxl

wb = openpyxl.load_workbook('DOCUMENTOS BASE/COSTOS_GASTOS_SEPTIEMBRE_2026.xlsx', data_only=True)

diff_target = 235920150 - 211683221.42
print(f"Diferencia exacta a explicar (235,920,150 - 211,683,221.42): {diff_target:,.2f}")

diff_with_inv = 235920150 - 223927171.42
print(f"Diferencia con inventario inicial incluido (235,920,150 - 223,927,171.42): {diff_with_inv:,.2f}")

# Let's check single transactions close to 24,236,928.58 or 11,992,978.58
ws = wb['Detalle Compras Septiembre 2026']
print("\n--- Búsqueda de valores cercanos a la diferencia en Detalle Compras ---")
for r in range(2, ws.max_row + 1):
    vr = float(ws.cell(r, 7).value or 0)
    if abs(vr - diff_target) < 1000000 or abs(vr - diff_with_inv) < 1000000:
        print(f"Fila {r}: {ws.cell(r, 3).value} | {ws.cell(r, 4).value} | Vr: {vr:,.2f} | Factura: {ws.cell(r, 8).value}")

# Check in 'Transacciones Repetidas'
ws_rep = wb['Transacciones Repetidas']
print("\n--- Búsqueda en Transacciones Repetidas ---")
for r in range(2, ws_rep.max_row + 1):
    vr = float(ws_rep.cell(r, 7).value or 0)
    if abs(vr - diff_target) < 1000000 or abs(vr - diff_with_inv) < 1000000:
        print(f"Fila Repetida {r}: {ws_rep.cell(r, 3).value} | {ws_rep.cell(r, 4).value} | Vr: {vr:,.2f}")

# Check groups by provider in 'Facturas por Proveedor'
ws_prov = wb['Facturas por Proveedor']
print("\n--- Proveedores con saldo o facturación cercana a la diferencia ---")
for r in range(11, ws_prov.max_row + 1):
    prov = ws_prov.cell(r, 2).value
    fact = float(ws_prov.cell(r, 5).value or 0)
    pagado = float(ws_prov.cell(r, 6).value or 0)
    saldo = float(ws_prov.cell(r, 7).value or 0)
    if abs(fact - diff_target) < 2000000 or abs(saldo - diff_target) < 2000000 or abs(pagado - diff_target) < 2000000:
        print(f"Prov: {prov} | Facturado: {fact:,.2f} | Pagado: {pagado:,.2f} | Saldo: {saldo:,.2f}")
